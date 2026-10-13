import openpyxl
import argparse
from pathlib import Path

parser = argparse.ArgumentParser(description="Genera un SQL para importar un catálogo a una base vacía.")
parser.add_argument("--bandas", required=True, type=Path, help="Excel del catálogo de bandas")
parser.add_argument("--rulimanes", required=True, type=Path, help="Excel del catálogo de rulimanes")
parser.add_argument("--orings", required=True, type=Path, help="Excel del catálogo de orings")
parser.add_argument("--salida", type=Path, help="Ruta del SQL generado (por defecto, junto a este archivo)")
args = parser.parse_args()

SOURCES = {
    "bandas": args.bandas,
    "rulimanes": args.rulimanes,
    "orings": args.orings,
}
OUTPUT = args.salida or Path(__file__).with_name("importar_catalogo_actual.sql")


def text(value):
    if value is None:
        return None
    value = str(value).strip()
    return value or None


def sql_value(value):
    if value is None or value == "":
        return "NULL"
    if isinstance(value, (int, float)):
        return str(value)
    return "'" + str(value).replace("'", "''") + "'"


def rows(sheet):
    values = list(sheet.iter_rows(min_row=2, values_only=True))
    return [row for row in values if any(value is not None and str(value).strip() for value in row)]


catalog = []
bandas = openpyxl.load_workbook(SOURCES["bandas"], read_only=True, data_only=True)

for sheet_name, subtype, uses_description in [
    ("LOCALPK", "Canales", True),
    ("BANDA LISAS A-B-HM-3L-4L-5L", "Lisas", False),
    ("DENTADA A-B-HM", "Dentadas", False),
    ("BANDAS DISTRIBUCION", "Distribución", True),
]:
    for row in rows(bandas[sheet_name]):
        if sheet_name == "BANDAS DISTRIBUCION":
            product_id, code, description, price, stock, brand = row
        elif uses_description:
            product_id, code, description, price, brand, stock = row
        else:
            product_id, code, price, stock, brand = row
            description = None
        code = text(code)
        catalog.append((product_id, f"Banda {code}", code, text(description), subtype,
                        None, text(description) if subtype == "Canales" else None,
                        text(brand), price or 0, stock or 0, "Bandas"))

rulimanes = openpyxl.load_workbook(SOURCES["rulimanes"], read_only=True, data_only=True).active
for product_id, code, stock, brand, price in rows(rulimanes):
    code = text(code)
    catalog.append((product_id, f"Ruliman {code}", code, None, None, None, None,
                    text(brand), price or 0, stock or 0, "Rulimanes"))

orings = openpyxl.load_workbook(SOURCES["orings"], read_only=True, data_only=True).active
for product_id, code, price_nitrilo, price_viton, stock_nitrilo, stock_viton in rows(orings):
    code = text(code)
    for material, price, stock, offset in [
        ("Nitrilo", price_nitrilo, stock_nitrilo, 0),
        ("Viton", price_viton, stock_viton, 1),
    ]:
        catalog.append((product_id * 10 + offset, f"Oring {code} - {material}", code,
                        None, None, material, None, None, price or 0, stock or 0, "Orings"))

ids = [row[0] for row in catalog]
if len(ids) != len(set(ids)):
    raise ValueError("Los identificadores generados para producto no son únicos.")

with OUTPUT.open("w", encoding="utf-8", newline="\n") as file:
    file.write("-- Generado desde los tres archivos de inventario proporcionados.\n")
    file.write("-- Solo importa en una base vacía. No borra ventas, productos ni movimientos.\n")
    file.write("USE inventario_db;\n\n")
    file.write("DROP PROCEDURE IF EXISTS validar_importacion_catalogo;\n")
    file.write("DELIMITER $$\n")
    file.write("CREATE PROCEDURE validar_importacion_catalogo()\nBEGIN\n")
    file.write("  IF EXISTS (SELECT 1 FROM producto LIMIT 1) OR EXISTS (SELECT 1 FROM venta LIMIT 1) OR EXISTS (SELECT 1 FROM movimiento LIMIT 1) THEN\n")
    file.write("    SIGNAL SQLSTATE '45000' SET MESSAGE_TEXT = 'Importación cancelada: la base contiene productos, ventas o movimientos.';\n")
    file.write("  END IF;\nEND$$\nDELIMITER ;\n")
    file.write("CALL validar_importacion_catalogo();\nDROP PROCEDURE validar_importacion_catalogo;\n\n")
    file.write("INSERT INTO proveedor (nombre_proveedor, estado)\n")
    file.write("SELECT 'Proveedor por asignar', TRUE\n")
    file.write("WHERE NOT EXISTS (SELECT 1 FROM proveedor);\n")
    file.write("SET @proveedor_id = (SELECT MIN(id_proveedor) FROM proveedor);\n\n")
    for category in ("Bandas", "Rulimanes", "Orings"):
        file.write("INSERT INTO categoria (nombre_categoria)\n")
        file.write(f"SELECT '{category}' WHERE NOT EXISTS (SELECT 1 FROM categoria WHERE nombre_categoria = '{category}');\n")
    file.write("\n")
    file.write("INSERT INTO producto\n")
    file.write("(id_producto, nombre_producto, codigo, descripcion, tipo_banda, tipo_oring, aplicacion_vehiculo, marca, precio_producto, stock_producto, ubicacion_percha, id_categoria, id_proveedor)\nVALUES\n")
    statements = []
    for product_id, name, code, description, band_type, oring_type, application, brand, price, stock, category in catalog:
        category_lookup = f"(SELECT id_categoria FROM categoria WHERE nombre_categoria = '{category}')"
        statements.append(
            "(" + ", ".join([
                sql_value(product_id), sql_value(name), sql_value(code), sql_value(description),
                sql_value(band_type), sql_value(oring_type), sql_value(application), sql_value(brand),
                sql_value(round(float(price), 2)), sql_value(int(stock)), "NULL", category_lookup, "@proveedor_id",
            ]) + ")"
        )
    file.write(",\n".join(statements))
    file.write(";\n")

print(f"Generados {len(catalog)} productos en {OUTPUT}")
